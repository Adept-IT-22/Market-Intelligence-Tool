from docx import Document
import os

TEMPLATE_DIR = "Backend/report_templates"
os.makedirs(TEMPLATE_DIR, exist_ok=True)

def create_sprint_template():
    doc = Document()
    
    # Title
    doc.add_heading('{{ project_name }} - Sprint Report', 0)
    
    # Metadata table or section
    doc.add_paragraph('Sprinting Period: {{ start_date }} to {{ end_date }}')
    doc.add_paragraph('Sprint Number: {{ sprint_no }}')
    
    # Overall Status
    doc.add_heading('1. Executive Summary', level=1)
    doc.add_paragraph('{{ status }}')
    
    # Accomplishments
    doc.add_heading('2. Key Accomplishments', level=1)
    doc.add_paragraph('{{ progress }}')
    
    # Issues
    doc.add_heading('3. Issues & Blockers', level=1)
    doc.add_paragraph('{{ issues }}')
    
    # Next Steps
    doc.add_heading('4. Next Steps', level=1)
    doc.add_paragraph('{{ next_steps }}')
    
    path = os.path.join(TEMPLATE_DIR, "sprint_report.docx")
    doc.save(path)
    print(f"Template created at: {path}")

if __name__ == "__main__":
    create_sprint_template()
