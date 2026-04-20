from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os

def create_branded_template(report_key, title):
    doc = Document()
    
    # --- 1. Page Setup (Margins for Sidebar) ---
    section = doc.sections[0]
    section.left_margin = Inches(1.5) # Leave space for the blue sidebar bar
    section.right_margin = Inches(1.0)
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(1.0)

    # --- 2. Header (Logo) ---
    header = section.header
    header_para = header.paragraphs[0]
    
    logo_path = os.path.join("report_templates", "assets", "adept_logo.jpg")
    if os.path.exists(logo_path):
        run = header_para.add_run()
        run.add_picture(logo_path, width=Inches(1.8))
    else:
        header_para.text = "ADEPT TECHNOLOGIES"
    
    header_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    # --- 3. Sidebar (Simulation) ---
    # In python-docx, adding a true floating vertical bar on every page is tricky without lower-level XML.
    # For now, we will add a bold title and clear branding.

    # --- 4. Title & Metadata ---
    doc.add_paragraph("\n")
    report_title = doc.add_paragraph(title)
    report_title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = report_title.runs[0]
    run.font.size = Pt(26)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0x1D, 0x4E, 0x89) # Adept Blue

    # Metadata Table
    table = doc.add_table(rows=3, cols=2)
    table.style = 'Table Grid'
    
    # We use tags that docxtpl will fill
    cells = table.rows[0].cells
    cells[0].text = "Report Type:"
    cells[1].text = title
    
    cells = table.rows[1].cells
    cells[0].text = "Project/Campaign:"
    cells[1].text = "{{ project_name if project_name else campaign_name }}"
    
    cells = table.rows[2].cells
    cells[0].text = "Date:"
    cells[1].text = "{{ date if date else currentDate }}"

    doc.add_paragraph("\n")

    # --- 5. Sections Loop ---
    # This loop is for docxtpl to iterate over the sections dict
    doc.add_paragraph("{% for id, section in sections.items() %}")
    
    # Section Title
    sec_title = doc.add_paragraph("{{ section.title }}")
    sec_title.style = 'Heading 1'
    run = sec_title.runs[0]
    run.font.color.rgb = RGBColor(0x1D, 0x4E, 0x89)
    
    # Section Content
    doc.add_paragraph("{{ section.aiDraft }}")
    
    doc.add_paragraph("\n")
    doc.add_paragraph("{% endfor %}")

    # --- 6. Footer (Branded Contact) ---
    footer = section.footer
    footer_para = footer.paragraphs[0]
    footer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    footer_img_path = os.path.join("report_templates", "assets", "adept_footer.png")
    if os.path.exists(footer_img_path):
        run = footer_para.add_run()
        run.add_picture(footer_img_path, width=Inches(6.0))
    else:
        footer_para.text = "Adept Technologies Ltd — contact.adept@adept-techno.com"

    # Save
    filename = f"{report_key}_report.docx"
    output_path = os.path.join("report_templates", filename)
    doc.save(output_path)
    print(f"Template saved: {output_path}")

if __name__ == "__main__":
    templates = {
        "sprint": "Sprint Report",
        "marketing": "Marketing Performance Report",
        "weekly": "Weekly Activity Report",
        "monthly": "Monthly Strategic Overview"
    }
    
    for key, name in templates.items():
        create_branded_template(key, name)
